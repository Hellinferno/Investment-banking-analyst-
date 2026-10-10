"""Evidence exports with safe text, references and unique run filenames."""
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

from config import OUTPUT_ROOT


def sheet_text(value):
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def export_research(
    company,
    run_id,
    bundle,
    report,
    *,
    diligence=False,
    output_dir=OUTPUT_ROOT,
    review_items=None,
    export_id=None,
    include_xlsx=None,
):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^a-zA-Z0-9_-]", "_", company)[:70] or "Company"
    suffix = f"{run_id}_review_{export_id}" if export_id else run_id
    stem = output_dir / f"{safe_name}_{'Diligence' if diligence else 'Research'}_{suffix}"
    exported_at = datetime.now(timezone.utc).isoformat()
    snapshot = {
        "research_run_id": run_id,
        "exported_at": exported_at,
        "evidence": bundle,
        "report": report,
        "review_items": review_items or [],
    }
    json_path = stem.with_suffix(".json")
    json_path.write_text(json.dumps(snapshot, indent=2, ensure_ascii=False), encoding="utf-8")
    pdf_path = stem.with_suffix(".pdf")
    write_research_pdf(pdf_path, company, bundle, report, review_items=review_items or [], run_id=run_id, exported_at=exported_at, diligence=diligence)
    outputs = [(str(json_path), "json"), (str(pdf_path), "pdf")]
    if include_xlsx is None:
        include_xlsx = diligence
    if include_xlsx:
        xlsx_path = stem.with_suffix(".xlsx")
        write_diligence_workbook(xlsx_path, bundle, report)
        outputs.append((str(xlsx_path), "xlsx"))
    return outputs


def write_research_pdf(path, company, bundle, report, *, review_items=None, run_id="", exported_at="", diligence=False):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    styles = getSampleStyleSheet()
    font_dir = Path(os.getenv("AIBAA_PDF_FONT_DIR", "/usr/share/fonts/truetype/dejavu"))
    if (font_dir / "DejaVuSans.ttf").exists() and (font_dir / "DejaVuSans-Bold.ttf").exists():
        pdfmetrics.registerFont(TTFont("ResearchSans", str(font_dir / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("ResearchSans-Bold", str(font_dir / "DejaVuSans-Bold.ttf")))
        pdfmetrics.registerFontFamily("ResearchSans", normal="ResearchSans", bold="ResearchSans-Bold",
                                     italic="ResearchSans", boldItalic="ResearchSans-Bold")
        for name in ("BodyText", "Title", "Heading1", "Heading2", "Heading3"):
            styles[name].fontName = "ResearchSans" if name == "BodyText" else "ResearchSans-Bold"
    styles.add(ParagraphStyle("Evidence", parent=styles["BodyText"], fontSize=9, leading=13, spaceAfter=7))
    styles.add(ParagraphStyle("Note", parent=styles["Evidence"], textColor=colors.HexColor("#555555"), fontSize=8))
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm, title=f"{company} - {report['title']}")
    def p(text, style="Evidence"):
        text = str(text)
        if styles[style].fontName.startswith("Helvetica"):
            text = text.replace("₹", "INR ")
        return Paragraph(escape(text), styles[style])
    story = [p("AIBAA | Public Diligence Discovery" if diligence else "AIBAA | Company Intelligence", "Title"), p(report["title"], "Heading1"),
             p(company, "Heading2"), p(f"Mode: {bundle['mode'].upper()} | Coverage: {bundle['status']} | Synthesis: {report['synthesis_mode']}", "Note"),
             p(f"Research run: {run_id or 'Not provided'} | Retrieved: {bundle['collected_at']}", "Note"),
             p(f"Exported: {exported_at}", "Note"), p(report["summary"])]
    for warning in report["warnings"]:
        story.append(p(warning, "Note"))
    story.extend([Spacer(1, 0.25*cm), p("Observations for review", "Heading2")])
    for finding in report["findings"]:
        story.append(p(f"[{', '.join(finding['source_ids'])}] {finding['category']}: {finding['statement']}"))
    if not report["findings"]:
        story.append(p("No supported observations available."))
    if review_items:
        story.append(p("Analyst Review Board", "Heading2"))
        story.append(p("Reviewed means handled by an analyst; it does not certify a search claim as true.", "Note"))
        for item in review_items:
            story.append(p(f"{item['kind'].title()} | {item['status'].replace('_', ' ').title()} | {item['title']}", "Heading3"))
            story.append(p(item.get("note") or "No analyst note provided."))
            if item.get("next_action"):
                story.append(p(f"Next action: {item['next_action']}"))
            story.append(p(f"Sources: {', '.join(item['source_ids'])} | Updated by: {item['updated_by']} | Updated: {item['updated_at']}", "Note"))
    story.append(p("Sources and provenance", "Heading2"))
    for source in bundle["sources"]:
        story.append(p(f"{source['id']} | {source['title']}", "Heading3"))
        url = escape(source["url"], {'"': '&quot;'})
        story.append(Paragraph(f'<link href="{url}" color="#225599">{url}</link>', styles["Note"]))
        story.append(p(f"Publisher: {source['publisher']} | Published: {source['published_at'] or 'Not provided'}", "Note"))
        story.append(p(f"Engine: {source['engine']} | Search ID: {source['search_id']} | Cached: {source['cache_hit']}", "Note"))
        story.append(p(f"Query: {source['query']}", "Note"))
    story.append(p("Search coverage", "Heading2"))
    for search in bundle["searches"]:
        story.append(p(f"{search['purpose']}: {search['status']} - {search['result_count']} results"))
        if search.get("error"):
            story.append(p(search["error"], "Note"))
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(2*cm, 1.2*cm, "AIBAA | Analyst review required")
        canvas.drawRightString(A4[0]-2*cm, 1.2*cm, f"Page {document.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def write_diligence_workbook(path, bundle, report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Company", sheet_text(bundle["company_name"])])
    summary.append(["Mode", bundle["mode"]])
    summary.append(["Coverage", bundle["status"]])
    summary.append(["Risk score", "Not assessed - discovery only"])
    summary.append(["Retrieved", bundle["collected_at"]])
    for warning in report["warnings"]:
        summary.append(["Review note", sheet_text(warning)])
    checklist = wb.create_sheet("Review Checklist")
    checklist.append(["Category", "Observation", "Sources", "Verification", "Status"])
    for finding in report["findings"]:
        checklist.append([sheet_text(finding["category"]), sheet_text(finding["statement"]),
                          ", ".join(finding["source_ids"]), finding["verification"], "Needs review"])
    sources = wb.create_sheet("Sources")
    sources.append(["ID", "Title", "URL", "Published", "Retrieved", "Engine", "Search ID", "Cached"])
    for source in bundle["sources"]:
        sources.append([source["id"], sheet_text(source["title"]), sheet_text(source["url"]),
                        sheet_text(source["published_at"]), source["retrieved_at"], source["engine"],
                        source["search_id"], str(source["cache_hit"])])
        url_cell = sources.cell(row=sources.max_row, column=3)
        url_cell.hyperlink = source["url"]
        url_cell.style = "Hyperlink"
    column_widths = {
        "Summary": [20, 90],
        "Review Checklist": [24, 90, 16, 24, 18],
        "Sources": [10, 48, 70, 18, 25, 16, 28, 10],
    }
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="16324F")
        for index, width in enumerate(column_widths[sheet.title], start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row[0].row].height = 60
    wb.save(path)
